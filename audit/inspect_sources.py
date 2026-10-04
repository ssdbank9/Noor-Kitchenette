import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

def normal_value(value):
    if isinstance(value, openpyxl.worksheet.formula.ArrayFormula):
        return {"kind": "array_formula", "text": value.text, "range": value.ref}
    return value

ROOT = Path(__file__).resolve().parent
SOURCES = [
    Path(r"C:\Users\Aly Jafferani\Desktop\Noor_Kitchen_Planner_v2.xlsx"),
    Path(r"C:\Users\Aly Jafferani\Documents\Codex\2026-07-08\fianc-video-file-project\outputs\Noor_Kitchen_Planner_Final.xlsx"),
]
result = []
for path in SOURCES:
    wb = openpyxl.load_workbook(path, data_only=False)
    cache = openpyxl.load_workbook(path, data_only=True)
    record = {"file": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
              "defined_names": [{"name": n, "text": v.attr_text} for n, v in wb.defined_names.items()],
              "sheets": []}
    for ws in wb.worksheets:
        cells = []
        for row in ws:
            for c in row:
                if c.value is not None or c.comment or c.hyperlink:
                    cells.append({"cell": c.coordinate, "value": normal_value(c.value),
                                  "type": c.data_type, "cached": cache[ws.title][c.coordinate].value,
                                  "comment": c.comment.text if c.comment else None,
                                  "link": c.hyperlink.target if c.hyperlink else None})
        validation = [{"type": dv.type, "range": str(dv.sqref), "formula1": dv.formula1,
                       "formula2": dv.formula2, "prompt": dv.prompt, "error": dv.error}
                      for dv in ws.data_validations.dataValidation]
        sheet = {"name": ws.title, "state": ws.sheet_state, "rows": ws.max_row, "columns": ws.max_column,
                 "cells": cells, "validation": validation,
                 "merged": [str(x) for x in ws.merged_cells.ranges],
                 "hidden_rows": [k for k, v in ws.row_dimensions.items() if v.hidden],
                 "hidden_columns": [k for k, v in ws.column_dimensions.items() if v.hidden],
                 "tables": list(ws.tables), "charts": len(ws._charts), "images": len(ws._images),
                 "conditional_rules": len(ws.conditional_formatting)}
        record["sheets"].append(sheet)
    result.append(record)
(ROOT / "workbook-evidence.json").write_text(json.dumps(result, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
for book in result:
    print("FILE", Path(book["file"]).name, "SHA256", book["sha256"])
    print("NAMES", json.dumps(book["defined_names"], ensure_ascii=False))
    for sheet in book["sheets"]:
        formulas = [c for c in sheet["cells"] if c["type"] == "f"]
        print(json.dumps({"sheet": sheet["name"], "state": sheet["state"],
                          "dimensions": [sheet["rows"], sheet["columns"]], "nonempty": len(sheet["cells"]),
                          "formulas": len(formulas), "validation": sheet["validation"],
                          "charts": sheet["charts"], "images": sheet["images"],
                          "preview": [{"cell": c["cell"], "value": c["value"]}
                                      for c in sheet["cells"][:14]]}, ensure_ascii=False, default=str))

html_path = Path(r"C:\Users\Aly Jafferani\Documents\Codex\2026-07-08\fianc-video-file-project\outputs\Noor_Kitchen_App_v3.html")
html = html_path.read_text(encoding="utf-8")
match = re.search(r"const DB = (\{.*\});", html)
if match:
    db = json.loads(match.group(1))
    (ROOT / "html-database.json").write_text(json.dumps(db, indent=2, ensure_ascii=False), encoding="utf-8")
    print("HTML_DATABASE", {k: len(v) if hasattr(v, "__len__") else type(v).__name__ for k, v in db.items()})
lines = html.splitlines()
(ROOT / "html-source-numbered.txt").write_text("\n".join(f"{i+1}: {line}" for i, line in enumerate(lines)), encoding="utf-8")
print("HTML", len(lines), "lines", "SHA256", hashlib.sha256(html_path.read_bytes()).hexdigest())
